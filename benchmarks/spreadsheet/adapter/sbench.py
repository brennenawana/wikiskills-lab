#!/usr/bin/env python3
"""Adapter: wires the engine to SpreadsheetBench Verified-400.

Task loading (the Verified-400 `{tc}_{id}_init.xlsx` / `_golden.xlsx`
naming), the upstream checker binding, and the agentic executor (one
Claude CLI conversation per test case, Bash tool only, isolated workdir,
blind to the wiki by construction).

Ground-truth isolation: only `compare()` in this module ever touches golden
files; executor workdirs receive input copies only.

Needs: `openpyxl` (and the checker needs `pandas`) at the upstream pins —
see README.md here. Data location: ./data by default, or $SB_DATA_DIR.
"""

import importlib.util
import json
import os
import shutil
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
REPO = HERE.parents[2]
sys.path.insert(0, str(REPO / "engine" / "gateway"))
from gateway import Gateway, GatewayError               # noqa: E402

DATA_DIR = Path(os.environ.get("SB_DATA_DIR", HERE.parent / "data"))
DATASET_DIR = DATA_DIR / "spreadsheetbench_verified_400"
UPSTREAM_DIR = DATA_DIR / "upstream"
PROMPT_FILE = REPO / "steps" / "5-improve" / "prompts" / \
    "inference-agent-spreadsheetbench.txt"

TURN_CAP = 15
CMD_TIMEOUT = 90

_eval_mod = None


# ------------------------------------------------------------------ tasks

def load_dataset():
    path = DATASET_DIR / "dataset.json"
    if not path.exists():
        raise SystemExit("Dataset not found at %s — run fetch_data.py first."
                         % path)
    return json.loads(path.read_text(encoding="utf-8"))


def load_manifest():
    return json.loads((HERE.parent / "manifest.json").read_text(
        encoding="utf-8"))


def tasks_by_id(ids):
    idx = {str(t["id"]): t for t in load_dataset()}
    return [idx[str(i)] for i in ids]


def task_files(task, tc):
    d = DATASET_DIR / "spreadsheet" / str(task["id"])
    return (d / ("%d_%s_init.xlsx" % (tc, task["id"])),
            d / ("%d_%s_golden.xlsx" % (tc, task["id"])))


def test_cases(task):
    out = []
    for tc in (1, 2, 3):
        i, g = task_files(task, tc)
        if i.exists() and g.exists():
            out.append(tc)
    return out


def spreadsheet_preview(path, max_rows=5, max_cols=10, max_chars=2000):
    import openpyxl
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        return "(could not read spreadsheet: %s)" % exc
    parts = []
    for name in wb.sheetnames:
        ws = wb[name]
        parts.append("[Sheet: %s] (%d rows x %d cols)"
                     % (name, ws.max_row, ws.max_column))
        for i, row in enumerate(ws.iter_rows(max_row=max_rows,
                                             max_col=max_cols,
                                             values_only=True)):
            parts.append("\t".join("" if v is None else str(v)[:40]
                                   for v in row))
            if i + 1 >= max_rows:
                break
    wb.close()
    text = "\n".join(parts)
    return text[:max_chars] + ("\n..." if len(text) > max_chars else "")


# ---------------------------------------------------------------- checker

def _upstream():
    global _eval_mod
    if _eval_mod is None:
        path = UPSTREAM_DIR / "evaluation.py"
        if not path.exists():
            raise SystemExit("Checker missing at %s — run fetch_data.py."
                             % path)
        recorded = (UPSTREAM_DIR / "COMMIT").read_text().strip()
        expected = "49b73a94775fb489063f60ca1865e3a650079a79"
        if recorded != expected:
            raise SystemExit("Checker commit mismatch; refusing to score.")
        spec = importlib.util.spec_from_file_location("sb_evaluation",
                                                      str(path))
        mod = importlib.util.module_from_spec(spec)
        sys.modules["sb_evaluation"] = mod
        spec.loader.exec_module(mod)
        _eval_mod = mod
    return _eval_mod


def compare(gt_file, proc_file, task):
    import contextlib
    import io
    try:
        with contextlib.redirect_stdout(io.StringIO()):
            ok, msg = _upstream().compare_workbooks(
                str(gt_file), str(proc_file),
                task["instruction_type"], task["answer_position"])
        return bool(ok), msg
    except Exception as exc:
        return False, "checker exception: %s" % exc


def score_task(task, output_paths):
    """Score = fraction of available test cases passed (the upstream
    'soft' metric; 395 of 400 tasks have exactly one test case)."""
    results = []
    for tc in test_cases(task):
        _, golden = task_files(task, tc)
        out = output_paths.get(tc)
        results.append(int(bool(out) and compare(golden, out, task)[0]))
    n = max(len(results), 1)
    return results.count(1) / n


def checker_selftest(task):
    """Suite eligibility: golden matches itself, input does not match
    golden, on every available test case."""
    tcs = test_cases(task)
    if not tcs:
        return False, "no init/golden files"
    for tc in tcs:
        inp, gold = task_files(task, tc)
        if not compare(gold, gold, task)[0]:
            return False, "golden fails its own check"
        if compare(gold, inp, task)[0]:
            return False, "input already passes"
    return True, ""


# --------------------------------------------------------------- executor

def _task_message(task, workdir, in_file, out_file):
    return ("working_directory: %s\ninstruction: %s\n"
            "spreadsheet_path: %s\nspreadsheet_content:\n%s\n"
            "instruction_type: %s\nanswer_position: %s\noutput_path: %s\n"
            % (workdir, task["instruction"], in_file,
               spreadsheet_preview(in_file), task["instruction_type"],
               task["answer_position"], out_file))


def _trace_from_events(events):
    lines = []
    for ev in events or []:
        t = ev.get("type")
        msg = ev.get("message") or {}
        for block in msg.get("content") or []:
            if not isinstance(block, dict):
                continue
            if t == "assistant" and block.get("type") == "text" \
                    and block.get("text"):
                lines.append("assistant: %s" % block["text"])
            elif t == "assistant" and block.get("type") == "tool_use":
                lines.append("[bash] %s"
                             % (block.get("input") or {}).get("command", ""))
            elif t == "user" and block.get("type") == "tool_result":
                content = block.get("content")
                if isinstance(content, list):
                    content = " ".join(str(c.get("text", ""))
                                       for c in content
                                       if isinstance(c, dict))
                lines.append("tool_result: %s" % str(content)[:1500])
    return "\n".join(lines)


def run_task(gateway, executor_cfg, task, skill_section, workroot, meter,
             run_tag):
    """One task = one CLI conversation per test case; returns the loop's
    result shape {"id", "score", "crash"} and writes traces on request via
    the returned "trace" text."""
    system = PROMPT_FILE.read_text(encoding="utf-8").replace(
        "{skill_section}", skill_section)
    outputs, traces, crash = {}, [], False
    for tc in test_cases(task):
        workdir = Path(workroot) / ("%s_tc%d" % (task["id"], tc))
        if workdir.exists():
            shutil.rmtree(workdir)
        workdir.mkdir(parents=True)
        src_in, _ = task_files(task, tc)
        in_file = workdir / src_in.name
        shutil.copy2(src_in, in_file)
        out_file = workdir / ("%d_%s_output.xlsx" % (tc, task["id"]))

        if meter is not None:
            meter.precheck(0.10, tags=run_tag)
        try:
            result = gateway.call(
                executor_cfg, _task_message(task, workdir, in_file,
                                            out_file),
                system=system,
                options={"tools": "Bash", "max_turns": TURN_CAP,
                         "cwd": str(workdir), "stream": True,
                         "timeout_s": TURN_CAP * CMD_TIMEOUT + 300,
                         "extra_env": {
                             "BASH_DEFAULT_TIMEOUT_MS":
                                 str(CMD_TIMEOUT * 1000),
                             "BASH_MAX_TIMEOUT_MS":
                                 str(CMD_TIMEOUT * 1000)}})
        except GatewayError as exc:
            crash = True
            traces.append("[harness crash: %s]" % exc)
            continue
        if meter is not None:
            cost = result["usd"] if meter.unit == "usd" else \
                (result["tokens_in"] + result["tokens_out"])
            meter.record(cost, tags=dict(run_tag, task=str(task["id"])),
                         model=result.get("model"))
        traces.append(_trace_from_events(result.get("events")))
        if out_file.exists():
            outputs[tc] = out_file

    score = score_task(task, outputs)
    trace_text = ("=== TASK %s ===\nINSTRUCTION: %s\nSCORE: %.3f\n%s"
                  % (task["id"], task["instruction"], score,
                     "\n---\n".join(traces)))
    for tc in test_cases(task):
        shutil.rmtree(Path(workroot) / ("%s_tc%d" % (task["id"], tc)),
                      ignore_errors=True)
    return {"id": str(task["id"]), "score": score, "crash": crash,
            "trace": trace_text}


def make_rollout_fn(gateway, executor_cfg, workroot, meter, run_tag):
    """The engine loop's rollout function for this benchmark."""
    def rollout(tasks, skill_section, iteration, trace_dir):
        results = []
        for task in tasks:
            r = run_task(gateway, executor_cfg, task, skill_section,
                         workroot, meter, run_tag)
            if trace_dir is not None:
                (Path(trace_dir) / ("%s.txt" % r["id"])).write_text(
                    r["trace"], encoding="utf-8")
            results.append({k: r[k] for k in ("id", "score", "crash")})
        return results
    return rollout
